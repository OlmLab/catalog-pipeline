"""Accession mining over cached Europe PMC JATS full text and supplementary ZIPs.
Reads harvest_lib's http_cache directly (no network). Produces mention dicts:
 paper_id, pmcid, accession, accession_type, source_section, context, from_supplement, file
"""
# --- microbiome_repo-pipeline repo layout shim (added 2026-09-26; original ran flat from one cwd) ---
import os as _os, sys as _sys
_here = _os.path.dirname(_os.path.abspath(__file__)) if "__file__" in globals() else _os.getcwd()
for _p in (_here, _os.path.join(_here, "..", "..")):
    _p = _os.path.abspath(_p)
    if _p not in _sys.path:
        _sys.path.insert(0, _p)
try:
    from catalog.models import resolve_model, set_host  # registers src/catalog/<stage>/ dirs on sys.path
except ImportError:  # flat-workspace mode (files copied side by side): models.py must sit alongside
    from models import resolve_model, set_host
set_host(globals().get("host"))  # kernel `host` is a frame global, not builtins (R1-01)
def _prompt_path(name):
    """cwd copy first (leaf-worker convention), else the packaged prompt under src/catalog/prompts/."""
    for _c in (name, _os.path.join("pilot_prompts", name), _os.path.join(_here, "..", "prompts", name)):
        if _os.path.exists(_c):
            return _c
    return name
# ---------------------------------------------------------------------------------------------
import io, json, os, re, sqlite3, zipfile
from pathlib import Path
from lxml import etree

ACC_PATTERNS = {
    "bioproject":  r"\bPRJ(?:NA|EB|DA|DB|CA)\d{3,9}\b",
    "sra_study":   r"\b(?:SRP|ERP|DRP)\d{5,9}\b",
    "sra_sample":  r"\b(?:SRS|ERS|DRS)\d{5,9}\b",
    "sra_experiment": r"\b(?:SRX|ERX|DRX)\d{5,9}\b",
    "sra_run":     r"\b(?:SRR|ERR|DRR)\d{5,9}\b",
    "biosample":   r"\bSAM(?:N|EA|EG|D)\d{5,12}\b",
    "geo":         r"\bGSE\d{3,7}\b",
    "ega":         r"\bEGA[SD]\d{8,13}\b",
    "dbgap":       r"\bphs\d{6}(?:\.v\d+\.p\d+)?\b",
    "gsa":         r"\b(?:CRA|HRA)\d{5,7}\b",
    "mgnify":      r"\bMGYS\d{8}\b",
}
COMPILED = {k: re.compile(v) for k, v in ACC_PATTERNS.items()}
# case-insensitive only for dbgap 'phs'
COMPILED["dbgap"] = re.compile(ACC_PATTERNS["dbgap"], re.I)
WIN = 200

CACHE_DB = Path("harvest_cache/http_cache.sqlite")
BLOB_DIR = Path("harvest_cache/blobs")

def cache_lookup(url):
    con = sqlite3.connect(CACHE_DB, timeout=60)
    row = con.execute("SELECT status, blob_sha, bytes FROM http_cache WHERE url=?", (url,)).fetchone()
    con.close()
    if not row or not row[1]:
        return None
    p = BLOB_DIR / row[1][:2] / row[1]
    if not p.exists():
        return None
    return {"status": str(row[0]), "bytes": row[2], "path": p}

def cache_index(prefix):
    con = sqlite3.connect(CACHE_DB, timeout=60)
    rows = con.execute("SELECT url, status, blob_sha, bytes FROM http_cache WHERE url LIKE ?", (prefix + "%",)).fetchall()
    con.close()
    return rows

def mine_text(text, section, paper_id, pmcid, from_supp=False, fname=None):
    out = []
    if not text:
        return out
    for kind, rx in COMPILED.items():
        for m in rx.finditer(text):
            acc = m.group(0)
            if kind != "dbgap":
                acc = acc.upper()
            else:
                acc = acc.lower()
            s, e = max(0, m.start() - WIN), min(len(text), m.end() + WIN)
            ctx = re.sub(r"\s+", " ", text[s:e]).strip()
            out.append({"paper_id": paper_id, "pmcid": pmcid, "accession": acc, "accession_type": kind,
                        "source_section": section, "context": ctx, "from_supplement": from_supp, "file": fname})
    return out

# ---------------------------------------------------------------- JATS
DA_TITLE = re.compile(r"data (?:and (?:code |materials? )?)?availability|availability of (?:supporting )?data|data access|accession|data deposition|data sharing|data statement|availability of data", re.I)
METH_TITLE = re.compile(r"method|materials|experimental procedure|study design|sequencing|sample collection|data collection|analysis|processing", re.I)
RES_TITLE = re.compile(r"result", re.I)

def _txt(el):
    return " ".join(etree.tostring(el, method="text", encoding="unicode").split())

def _sec_label(sec):
    st = (sec.get("sec-type") or "").lower()
    title_el = sec.find("title")
    title = _txt(title_el) if title_el is not None else ""
    if "data-availability" in st or "data availability" in st or DA_TITLE.search(title):
        return "data_availability"
    if "method" in st or "materials" in st or METH_TITLE.search(title):
        return "methods"
    if "result" in st or RES_TITLE.search(title):
        return "results"
    if "intro" in st or re.search(r"introduction|background", title, re.I):
        return "introduction"
    if "discussion" in st or re.search(r"discussion|conclusion", title, re.I):
        return "discussion"
    return None

def mine_jats(xml_bytes, paper_id, pmcid):
    out = []
    try:
        root = etree.fromstring(xml_bytes, parser=etree.XMLParser(recover=True, huge_tree=True))
    except Exception as e:
        return out, f"xml_parse_error: {e}"
    if root is None:
        return out, "xml_empty"
    # abstract
    for ab in root.iter("abstract"):
        out += mine_text(_txt(ab), "abstract", paper_id, pmcid)
    # data availability elements outside sec
    for el in root.iter("data-availability"):
        out += mine_text(_txt(el), "data_availability", paper_id, pmcid)
    for fn in root.iter("fn"):
        if (fn.get("fn-type") or "") in ("data-availability", "supplementary-material") or DA_TITLE.search(_txt(fn)[:120] or ""):
            out += mine_text(_txt(fn), "data_availability", paper_id, pmcid)
    # sections: label top-level and inherit downward
    def walk(sec, inherited):
        lab = _sec_label(sec) or inherited
        if lab is None and sec.find("title") is None:
            t0 = _txt(sec)[:100]
            if DA_TITLE.search(t0):
                lab = "data_availability"
        own_parts = []
        for child in sec:
            if child.tag == "sec":
                walk(child, lab)
            elif child.tag in ("title",):
                continue
            else:
                own_parts.append(_txt(child))
        text = " ".join(own_parts)
        out.extend(mine_text(text, lab or "body_other", paper_id, pmcid))
    body = root.find(".//body")
    if body is not None:
        for child in body:
            if child.tag == "sec":
                walk(child, None)
            else:
                out += mine_text(_txt(child), "body_other", paper_id, pmcid)
    # back matter (excluding ref-list): notes, sec, supplementary captions
    back = root.find(".//back")
    if back is not None:
        for child in back:
            if child.tag in ("ref-list",):
                continue
            if child.tag == "sec":
                walk(child, None)
            else:
                t = _txt(child)
                lab = "data_availability" if (DA_TITLE.search(t[:100]) or (child.tag == "notes" and (child.get("notes-type") or "").startswith("data"))) else "back_matter"
                out += mine_text(t, lab, paper_id, pmcid)
    # supplementary-material captions anywhere
    for sm in root.iter("supplementary-material"):
        out += mine_text(_txt(sm), "supplement_caption", paper_id, pmcid)
    # dedupe exact duplicates (same acc, section, context)
    seen, ded = set(), []
    for m in out:
        k = (m["accession"], m["source_section"], m["context"][:120])
        if k not in seen:
            seen.add(k); ded.append(m)
    return ded, None

# ---------------------------------------------------------------- supplements
TEXT_EXT = (".csv", ".tsv", ".txt", ".xml", ".json", ".md", ".fasta", ".fa", ".tab", ".html", ".htm")
MAX_MEMBER = 60 * 1024 * 1024

def _xlsx_text(b, limit_cells=400000):
    import openpyxl
    parts = []
    wb = openpyxl.load_workbook(io.BytesIO(b), read_only=True, data_only=True)
    n = 0
    for ws in wb.worksheets:
        parts.append(f"[sheet {ws.title}]")
        for row in ws.iter_rows(values_only=True):
            vals = [str(v) for v in row if v is not None]
            if vals:
                parts.append(" ".join(vals)); n += len(vals)
            if n > limit_cells:
                break
    return "\n".join(parts)

def _docx_text(b):
    with zipfile.ZipFile(io.BytesIO(b)) as z:
        names = [n for n in z.namelist() if n.startswith("word/") and n.endswith(".xml")]
        return " ".join(re.sub(r"<[^>]+>", " ", z.read(n).decode("utf8", "ignore")) for n in names)

def _pdf_text(b, max_pages=60):
    try:
        import pypdfium2 as pdfium
    except ImportError:
        return None
    pdf = pdfium.PdfDocument(io.BytesIO(b))
    parts = []
    for i in range(min(len(pdf), max_pages)):
        parts.append(pdf[i].get_textpage().get_text_range())
    return "\n".join(parts)

def mine_zip(zip_bytes, paper_id, pmcid):
    out, log = [], []
    try:
        z = zipfile.ZipFile(io.BytesIO(zip_bytes))
    except Exception as e:
        return out, [{"paper_id": paper_id, "file": None, "status": f"bad_zip: {e}"}]
    for info in z.infolist():
        name = info.filename
        low = name.lower()
        if info.is_dir():
            continue
        if info.file_size > MAX_MEMBER:
            log.append({"paper_id": paper_id, "file": name, "status": "skipped_gt60mb"}); continue
        try:
            b = z.read(info)
        except Exception as e:
            log.append({"paper_id": paper_id, "file": name, "status": f"read_error: {e}"}); continue
        text = None; status = "parsed"
        try:
            if low.endswith(TEXT_EXT):
                text = b.decode("utf8", "ignore")
            elif low.endswith((".xlsx", ".xlsm")):
                text = _xlsx_text(b)
            elif low.endswith(".xls"):
                status = "unparsed_xls"
            elif low.endswith(".docx"):
                text = _docx_text(b)
            elif low.endswith(".pdf"):
                text = _pdf_text(b)
                if text is None:
                    status = "unparsed_pdf_no_pypdfium2"
            elif low.endswith(".zip"):
                sub, sublog = mine_zip(b, paper_id, pmcid)
                for m in sub:
                    m["file"] = name + "/" + (m["file"] or "")
                out += sub; log += sublog; continue
            else:
                status = "skipped_binary"
        except Exception as e:
            status = f"parse_error: {type(e).__name__}"
        if text:
            out += mine_text(text, f"supplement:{name}", paper_id, pmcid, from_supp=True, fname=name)
        log.append({"paper_id": paper_id, "file": name, "status": status, "bytes": info.file_size})
    return out, log
